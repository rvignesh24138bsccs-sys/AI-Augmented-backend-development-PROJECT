import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = Document()

# Set standard margins
sections = doc.sections
for section in sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Read Markdown file
with open(r"C:\fit track\PROJECT_DOCUMENTATION.md", "r", encoding="utf-8") as f:
    lines = f.readlines()

in_code_block = False
code_buffer = []

for line in lines:
    raw = line.rstrip("\r\n")

    if raw.startswith("```"):
        if in_code_block:
            # End code block
            code_text = "\n".join(code_buffer)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(code_text)
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(40, 40, 40)
            code_buffer = []
            in_code_block = False
        else:
            in_code_block = True
            code_buffer = []
        continue

    if in_code_block:
        code_buffer.append(raw)
        continue

    if not raw.strip():
        continue

    if raw.startswith("# "):
        title = raw[2:].strip()
        p = doc.add_heading(title, level=0)
        p.paragraph_format.space_after = Pt(12)
    elif raw.startswith("## "):
        title = raw[3:].strip()
        p = doc.add_heading(title, level=1)
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
    elif raw.startswith("### "):
        title = raw[4:].strip()
        p = doc.add_heading(title, level=2)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
    elif raw.startswith("#### "):
        title = raw[5:].strip()
        p = doc.add_heading(title, level=3)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
    elif raw.startswith("- ") or raw.startswith("* "):
        text = raw[2:].strip()
        p = doc.add_paragraph(style='List Bullet')
        # Clean basic bold markers
        parts = re.split(r'(\*\*[^*]+\*\*)', text)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
            else:
                p.add_run(part)
    elif raw.startswith("> "):
        text = raw[2:].strip()
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.4)
        r = p.add_run(text)
        r.italic = True
        r.font.color.rgb = RGBColor(80, 80, 80)
    elif raw.startswith("|"):
        # Table rows will be formatted as clean text
        if "---" in raw:
            continue
        cells = [c.strip() for c in raw.split("|")[1:-1]]
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("  •  ".join(cells))
        r.font.size = Pt(10)
    else:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        parts = re.split(r'(\*\*[^*]+\*\*)', raw)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
            else:
                p.add_run(part)

output_path = r"C:\Users\WIN\Downloads\FitTrack_AI_Project_Documentation.docx"
doc.save(output_path)
print(f"Word document saved to: {output_path}")
